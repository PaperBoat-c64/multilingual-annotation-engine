#!/usr/bin/env python3
"""
Output Writer — Multi-format export for annotated lyrics.

Generates polished HTML + DOCX by default.
Markdown is on-demand only.

Usage:
    python output_writer.py render --format html --data <json> --output <path>
    python output_writer.py render --format docx --data <json> --output <path>
    python output_writer.py render --format markdown --data <json> --output <path>
    python output_writer.py render-all --data <json> --output-dir <dir>

Input JSON format:
{
    "metadata": {
        "source_language": "ja",
        "target_language": "en",
        "title": "Song Title",
        "artist": "Artist Name"
    },
    "structure": {                        // from lyrics_reconstructor.py
        "structure_type": "verse-chorus",
        "sections": [
            {"type": "verse", "number": 1, "line_start": 1, "line_end": 4},
            {"type": "chorus", "number": 1, "line_start": 5, "line_end": 8}
        ]
    },
    "rows": [
        {
            "type": "original",           // one of: original, romanization, translation
            "language": "ja",
            "direction": "ltr",
            "segments": [
                {"text": "君", "color": "#E74C3C", "group_id": "g1",
                 "pronunciation": "kimi", "pos": "pronoun"}
            ]
        },
        ...
    ],
    "annotations": [
        {"group_id": "g5", "confidence": 0.75, "note": "Possible ambiguity..."}
    ],
    "vocabulary": [                        // optional — pre-computed vocabulary list
        {"word": "君", "romanization": "kimi", "color": "#E74C3C",
         "meaning": "you", "frequency": 8, "pos": "pronoun"}
    ]
}
"""

import argparse
import json
import sys
from pathlib import Path
from typing import Optional


def validate_data(data: dict) -> list[str]:
    """Validate the input data structure. Returns list of warnings."""
    warnings = []
    if "rows" not in data:
        raise ValueError("Input data must contain 'rows' key")
    if not isinstance(data["rows"], list):
        raise ValueError("'rows' must be a list")
    for i, row in enumerate(data["rows"]):
        if "segments" not in row:
            warnings.append(f"Row {i} has no segments")
        for j, seg in enumerate(row.get("segments", [])):
            if "text" not in seg:
                warnings.append(f"Row {i}, segment {j}: missing 'text'")
            if "color" not in seg:
                warnings.append(f"Row {i}, segment {j}: missing 'color'")

    return warnings


def escape_html(text: str) -> str:
    """Escape HTML special characters."""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def escape_xml(text: str) -> str:
    """Escape XML special characters."""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


# Tooltip lookup and HTML generation are now handled client-side
# by the template's JavaScript (assets/templates/lyrics.html).
# The Python script embeds the raw annotation JSON; JS builds its own lookup.


def render_html(data: dict) -> str:
    """
    Render annotated data to HTML using the Apple Music template.

    Reads the template from assets/templates/lyrics.html, injects the
    annotation JSON data via placeholder replacement, and returns the
    complete HTML document. All CSS, JS, and DOM construction is handled
    client-side by the template.

    Falls back to a minimal embedded renderer if the template file is missing.
    """
    template_path = Path(__file__).resolve().parent.parent / "assets" / "templates" / "lyrics.html"

    if not template_path.exists():
        print(f"Warning: Template not found at {template_path}, using minimal fallback", file=sys.stderr)
        return _render_html_minimal(data)

    template = template_path.read_text(encoding="utf-8")

    metadata = data.get("metadata", {})
    title = escape_html(metadata.get("title", "Annotated Document"))
    source_lang = escape_html(metadata.get("source_language", ""))
    target_lang = escape_html(metadata.get("target_language", ""))

    metadata_json = json.dumps(metadata, ensure_ascii=False)
    rows_json = json.dumps(data.get("rows", []), ensure_ascii=False)
    vocabulary_json = json.dumps(data.get("vocabulary", []), ensure_ascii=False)
    annotations_json = json.dumps(data.get("annotations", []), ensure_ascii=False)
    structure_json = json.dumps(
        data.get("structure", metadata.get("structure", {})), ensure_ascii=False
    )

    html = template.replace("{{TITLE}}", title)
    html = html.replace("{{SOURCE_LANG}}", source_lang)
    html = html.replace("{{TARGET_LANG}}", target_lang)
    html = html.replace("{{METADATA_JSON}}", metadata_json)
    html = html.replace("{{ROWS_JSON}}", rows_json)
    html = html.replace("{{VOCABULARY_JSON}}", vocabulary_json)
    html = html.replace("{{ANNOTATIONS_JSON}}", annotations_json)
    html = html.replace("{{STRUCTURE_JSON}}", structure_json)

    return html


def _render_html_minimal(data: dict) -> str:
    """
    Fallback: generate minimal HTML if the template file is missing.
    Preserves basic readability without the full Apple Music design system.
    """
    title = escape_html(data.get("metadata", {}).get("title", "Annotated Document"))
    metadata_json = json.dumps(data.get("metadata", {}), ensure_ascii=False)
    rows_json = json.dumps(data.get("rows", []), ensure_ascii=False)
    vocabulary_json = json.dumps(data.get("vocabulary", []), ensure_ascii=False)
    annotations_json = json.dumps(data.get("annotations", []), ensure_ascii=False)
    structure_json = json.dumps(
        data.get("structure", data.get("metadata", {}).get("structure", {})), ensure_ascii=False
    )

    return f"""<!DOCTYPE html>
<html><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1.0">
<title>{title}</title>
<style>
body{{font-family:-apple-system,BlinkMacSystemFont,sans-serif;background:#000;color:#fff;
max-width:680px;margin:0 auto;padding:80px 24px;line-height:1.6}}
h1{{text-align:center;font-size:28px;font-weight:700;margin-bottom:8px}}
.meta{{text-align:center;font-size:15px;color:rgba(255,255,255,.5);margin-bottom:80px}}
.line-group{{margin-bottom:48px;text-align:center}}
.row.original{{font-size:24px;font-weight:600;margin-bottom:8px}}
.row.romanization{{font-size:15px;font-style:italic;color:rgba(255,255,255,.45);margin-bottom:6px}}
.row.translation{{font-size:18px;font-weight:500}}
.seg{{display:inline-block;padding:2px 6px;margin:1px 3px;border-radius:4px}}
.warning{{text-align:center;color:rgba(255,200,40,.7);font-size:14px;margin-bottom:60px;
border:1px solid rgba(255,200,40,.2);padding:16px;border-radius:8px}}
</style></head><body>
<h1>{title}</h1>
<div class="warning">⚠ Template file not found. Reinstall the skill to restore full Apple Music design.</div>
<div id="app"><p style="text-align:center;color:rgba(255,255,255,.3)">Loading data…</p></div>
<script>
window.__ANNOTATION_DATA__={{{{metadata:{metadata_json},rows:{rows_json},
vocabulary:{vocabulary_json},annotations:{annotations_json},structure:{structure_json}}}}};
(function(){{var D=window.__ANNOTATION_DATA__;var app=document.getElementById('app');
app.innerHTML='';var h='';var rows=D.rows||[];var lg=[];var cg=[];
for(var i=0;i<rows.length;i++){{if(rows[i].type==='original'&&cg.length){{lg.push(cg);cg=[]}}cg.push(rows[i])}}
if(cg.length)lg.push(cg);if(!lg.length&&rows.length)lg=[rows];
for(var gi=0;gi<lg.length;gi++){{var g=lg[gi];h+='<div class="line-group">';
for(var ri=0;ri<g.length;ri++){{var row=g[ri];var segs=row.segments||[];
if(!segs.length)continue;var dir=row.direction==='rtl'?' dir="rtl"':'';
h+='<div class="row '+row.type+'"'+dir+'>';
for(var si=0;si<segs.length;si++){{var s=segs[si];
h+='<span class="seg" style="color:'+(s.color||'#fff')+'">'+(s.text||'')+'</span>'}}
h+='</div>'}}h+='</div>'}}app.innerHTML=h}})();
</script></body></html>"""


def render_markdown(data: dict) -> str:
    """
    Render annotated data to Markdown with custom color tags.
    NOTE: Markdown is on-demand only — not generated by default.
    """
    rows = data.get("rows", [])
    annotations = data.get("annotations", [])
    title = data.get("metadata", {}).get("title", "Annotated Document")
    structure = data.get("structure")

    md_parts = [f"# {title}", ""]
    line_index = 1

    # Build section map
    line_section_map = {}
    if structure and structure.get("sections"):
        for sec in structure["sections"]:
            label = sec.get("type", "").title()
            num = sec.get("number")
            section_name = f"{label} {num}" if num else label
            if "lines" in sec:
                for ln in sec["lines"]:
                    line_section_map[ln] = section_name
            elif "line_start" in sec and "line_end" in sec:
                for ln in range(sec["line_start"], sec["line_end"] + 1):
                    line_section_map[ln] = section_name

    # Group rows into line groups (same logic as HTML)
    line_groups = []
    current_group = []
    for row in rows:
        row_type = row.get("type", "")
        if row_type == "original" and current_group:
            line_groups.append(current_group)
            current_group = []
        current_group.append(row)
    if current_group:
        line_groups.append(current_group)
    if not line_groups:
        line_groups = [rows]

    for group in line_groups:
        section_name = line_section_map.get(line_index)
        if section_name:
            md_parts.extend(["", f"### {section_name}", ""])

        for row in group:
            segments = row.get("segments", [])
            if not segments:
                continue
            row_type = row.get("type", "translation")

            line_parts = []
            for seg in segments:
                text = seg.get("text", "")
                color = seg.get("color", "#000000")
                line_parts.append(f'<span style="color:{color}">{text}</span>')

            prefix = {
                "original": "**原文** ",
                "romanization": "*音译* ",
                "translation": "> ",
                "ipa": "IPA ",
                "gloss": "注释 ",
                "notes": "📝 ",
            }.get(row_type, "")

            md_parts.append(f"{prefix}{' '.join(line_parts)}  ")
        md_parts.append("")
        line_index += 1

    if annotations:
        md_parts.extend(["## Annotations", ""])
        for ann in annotations:
            if ann.get("confidence", 1.0) < 0.9:
                md_parts.append(
                    f"- ⚠ {ann.get('note', '')} (confidence: {ann.get('confidence', 0):.0%})"
                )

    return "\n".join(md_parts)


def render_docx(data: dict) -> bytes:
    """
    Render annotated data to a publication-quality DOCX document.

    Features:
    - Centered titles and section headers
    - Professional margins and spacing
    - Visual hierarchy through font size and weight
    - RTL support for Arabic/Persian/Hebrew
    - Clean, polished layout
    """
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
    except ImportError:
        return _render_docx_minimal(data)

    doc = Document()
    title = data.get("metadata", {}).get("title", "Annotated Document")
    source_lang = data.get("metadata", {}).get("source_language", "")
    target_lang = data.get("metadata", {}).get("target_language", "")
    rows = data.get("rows", [])
    annotations = data.get("annotations", [])
    structure = data.get("structure")
    vocabulary = data.get("vocabulary", [])

    # ---- Page setup ----
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.6

    # ---- Title ----
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(24)
    title_para.paragraph_format.space_after = Pt(6)

    # ---- Metadata subtitle ----
    meta_parts = []
    if source_lang:
        meta_parts.append(f"Source: {source_lang}")
    if target_lang:
        meta_parts.append(f"Target: {target_lang}")
    if meta_parts:
        meta_para = doc.add_paragraph()
        meta_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_run = meta_para.add_run(" · ".join(meta_parts))
        meta_run.font.size = Pt(10)
        meta_run.font.color.rgb = RGBColor(128, 128, 128)
        meta_para.paragraph_format.space_after = Pt(24)

    # ---- Section map ----
    line_section_map = {}
    if structure and structure.get("sections"):
        for sec in structure["sections"]:
            label = sec.get("type", "").title()
            num = sec.get("number")
            section_name = f"{label} {num}" if num else label
            if "lines" in sec:
                for ln in sec["lines"]:
                    line_section_map[ln] = section_name
            elif "line_start" in sec and "line_end" in sec:
                for ln in range(sec["line_start"], sec["line_end"] + 1):
                    line_section_map[ln] = section_name

    # ---- Group rows into line groups ----
    line_groups = []
    current_group = []
    for row in rows:
        row_type = row.get("type", "")
        if row_type == "original" and current_group:
            line_groups.append(current_group)
            current_group = []
        current_group.append(row)
    if current_group:
        line_groups.append(current_group)
    if not line_groups:
        line_groups = [rows]

    # ---- Render line groups ----
    line_index = 1
    for group in line_groups:
        section_name = line_section_map.get(line_index)

        # Section header
        if section_name:
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_before = Pt(12)
            spacer.paragraph_format.space_after = Pt(0)

            sec_para = doc.add_paragraph()
            sec_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sec_run = sec_para.add_run(section_name)
            sec_run.bold = True
            sec_run.font.size = Pt(11)
            sec_run.font.color.rgb = RGBColor(120, 120, 120)
            sec_run.font.small_caps = True
            sec_para.paragraph_format.space_after = Pt(8)

        for row in group:
            row_type = row.get("type", "translation")
            direction = row.get("direction", "ltr")
            segments = row.get("segments", [])

            if not segments:
                continue

            para = doc.add_paragraph()

            # Row-level formatting
            if row_type == "original":
                para.paragraph_format.space_before = Pt(6)
            else:
                para.paragraph_format.space_before = Pt(1)
            para.paragraph_format.space_after = Pt(1)

            # RTL support
            if direction == "rtl":
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                pPr = para._element.get_or_add_pPr()
                bidi = pPr.makeelement(qn("w:bidi"), {})
                pPr.append(bidi)

            for i, seg in enumerate(segments):
                text = seg.get("text", "")
                color_hex = seg.get("color", "#000000").lstrip("#")

                run = para.add_run(text)

                # Apply color
                try:
                    r, g, b = int(color_hex[0:2], 16), int(color_hex[2:4], 16), int(color_hex[4:6], 16)
                    run.font.color.rgb = RGBColor(r, g, b)
                except (ValueError, IndexError):
                    pass

                # Font size and weight per row type
                if row_type == "original":
                    run.bold = True
                    run.font.size = Pt(16)
                elif row_type == "romanization":
                    run.italic = True
                    run.font.size = Pt(11)
                elif row_type == "translation":
                    run.font.size = Pt(13)
                elif row_type == "gloss":
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(100, 100, 100)
                elif row_type == "ipa":
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(120, 120, 120)
                elif row_type == "notes":
                    run.font.size = Pt(10)
                    run.italic = True
                    run.font.color.rgb = RGBColor(120, 120, 120)

                # Space between segments (not after last)
                if i < len(segments) - 1:
                    spacer_run = para.add_run(" ")
                    spacer_run.font.size = Pt(10)

        line_index += 1

    # ---- Annotations ----
    low_conf = [a for a in annotations if a.get("confidence", 1.0) < 0.9]
    if low_conf:
        doc.add_paragraph()  # spacer
        annot_heading = doc.add_paragraph()
        annot_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading_run = annot_heading.add_run("Annotations")
        heading_run.bold = True
        heading_run.font.size = Pt(14)
        annot_heading.paragraph_format.space_after = Pt(8)

        for ann in low_conf:
            note = ann.get("note", "")
            conf = ann.get("confidence", 0)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(f"⚠ {note} (confidence: {conf:.0%})")
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(140, 120, 40)
            run.italic = True

    # ---- Vocabulary appendix ----
    if vocabulary:
        doc.add_paragraph()  # spacer
        vocab_heading = doc.add_paragraph()
        vocab_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading_run = vocab_heading.add_run("Vocabulary")
        heading_run.bold = True
        heading_run.font.size = Pt(14)
        vocab_heading.paragraph_format.space_after = Pt(12)

        # Table
        table = doc.add_table(rows=1 + len(vocabulary), cols=4)
        table.style = "Light Grid Accent 1"

        # Header
        for i, header_text in enumerate(["Word", "Pronunciation", "Meaning", "Frequency"]):
            cell = table.rows[0].cells[i]
            cell.text = header_text
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True
                    run.font.size = Pt(10)

        # Data rows
        for row_idx, entry in enumerate(vocabulary):
            word = entry.get("word", "")
            roman = entry.get("romanization", "")
            meaning = entry.get("meaning", "")
            freq = str(entry.get("frequency", 1))

            for col_idx, val in enumerate([word, roman, meaning, freq]):
                cell = table.rows[row_idx + 1].cells[col_idx]
                cell.text = val
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(10)
                        # Apply color to word column
                        if col_idx == 0 and entry.get("color"):
                            try:
                                ch = entry["color"].lstrip("#")
                                r, g, b = int(ch[0:2], 16), int(ch[2:4], 16), int(ch[4:6], 16)
                                run.font.color.rgb = RGBColor(r, g, b)
                            except (ValueError, IndexError):
                                pass

    # Save to bytes
    import io
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _render_docx_minimal(data: dict) -> bytes:
    """
    Fallback: Generate a minimal valid DOCX without python-docx dependency.
    Creates a simple XML-based DOCX with colored text runs.
    """
    import io
    import zipfile

    rows = data.get("rows", [])
    title = data.get("metadata", {}).get("title", "Annotated Document")

    # Build document.xml body
    body_parts = []

    # Title
    body_parts.append(
        f'<w:p><w:pPr><w:jc w:val="center"/></w:pPr>'
        f'<w:r><w:rPr><w:b/><w:sz w:val="40"/></w:rPr>'
        f'<w:t>{escape_xml(title)}</w:t></w:r></w:p>'
    )

    for row in rows:
        direction = row.get("direction", "ltr")
        segments = row.get("segments", [])
        row_type = row.get("type", "translation")
        rtl_attr = ' w:bidi="1"' if direction == "rtl" else ""

        para_runs = []
        for i, seg in enumerate(segments):
            text = escape_xml(seg.get("text", ""))
            color = seg.get("color", "#000000").lstrip("#")

            # Font size and style per row type
            sz = "24"  # 12pt default
            bold = ""
            italic = ""
            if row_type == "original":
                sz = "32"  # 16pt
                bold = "<w:b/>"
            elif row_type == "romanization":
                sz = "22"  # 11pt
                italic = "<w:i/>"
            elif row_type == "translation":
                sz = "26"  # 13pt
            elif row_type in ("gloss", "ipa", "notes"):
                sz = "20"  # 10pt

            para_runs.append(
                f'<w:r><w:rPr><w:color w:val="{color}"/><w:sz w:val="{sz}"/>'
                f'{bold}{italic}</w:rPr>'
                f'<w:t xml:space="preserve">{text}</w:t></w:r>'
            )

            # Space between segments
            if i < len(segments) - 1:
                para_runs.append(
                    f'<w:r><w:rPr><w:sz w:val="20"/></w:rPr>'
                    f'<w:t xml:space="preserve"> </w:t></w:r>'
                )

        bidi_val_attr = ' w:val="1"' if direction == "rtl" else ""
        body_parts.append(
            f'<w:p{rtl_attr}><w:pPr><w:bidi{bidi_val_attr}/></w:pPr>'
            f'{"".join(para_runs)}</w:p>'
        )

    document_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    {"".join(body_parts)}
  </w:body>
</w:document>"""

    # Build minimal DOCX (Zip of XML files)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>""")
        zf.writestr("_rels/.rels", """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>""")
        zf.writestr("word/document.xml", document_xml)

    return buf.getvalue()


def main():
    parser = argparse.ArgumentParser(
        description="Output Writer for Multilingual Lyrics Translation Engine"
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    # Single-format render
    render_parser = subparsers.add_parser("render", help="Render to a single format")
    render_parser.add_argument("--format", required=True,
                               choices=["html", "docx", "markdown"],
                               help="Output format (html + docx are default)")
    render_parser.add_argument("--data", required=True, help="JSON string of annotated data")
    render_parser.add_argument("--output", required=True, help="Output file path")

    # Render all default formats
    render_all_parser = subparsers.add_parser("render-all", help="Render HTML + DOCX (all defaults)")
    render_all_parser.add_argument("--data", required=True, help="JSON string of annotated data")
    render_all_parser.add_argument("--output-dir", required=True, help="Output directory path")
    render_all_parser.add_argument("--stem", default="lyrics", help="Output file stem (default: lyrics)")

    args = parser.parse_args()

    try:
        data = json.loads(args.data)
    except json.JSONDecodeError as e:
        print(json.dumps({"error": f"Invalid JSON input: {e}"}), file=sys.stderr)
        sys.exit(1)

    try:
        warnings = validate_data(data)
        if warnings:
            for w in warnings:
                print(f"Warning: {w}", file=sys.stderr)

        if args.command == "render":
            result = _render_single(data, args.format, args.output)
            print(json.dumps(result, ensure_ascii=False))

        elif args.command == "render-all":
            output_dir = Path(args.output_dir)
            output_dir.mkdir(parents=True, exist_ok=True)
            stem = args.stem

            results = {}
            errors = {}
            formats = [
                ("html", f"{stem}.html"),
                ("docx", f"{stem}.docx"),
            ]

            for fmt, filename in formats:
                output_path = output_dir / filename
                try:
                    result = _render_single(data, fmt, str(output_path))
                    results[fmt] = result
                except Exception as e:
                    errors[fmt] = str(e)
                    print(f"Warning: {fmt} generation failed: {e}", file=sys.stderr)

            status = "partial" if errors else "ok"
            print(json.dumps({
                "status": status,
                "output_dir": str(output_dir.resolve()),
                "results": results,
                "errors": errors if errors else None,
            }, ensure_ascii=False))

    except Exception as e:
        print(json.dumps({"error": str(e)}), file=sys.stderr)
        sys.exit(1)


def _render_single(data: dict, fmt: str, output: str) -> dict:
    """Render data to a single format. Returns result dict."""
    if fmt == "html":
        content = render_html(data)
        mode = "w"
        encoding = "utf-8"
    elif fmt == "markdown":
        content = render_markdown(data)
        mode = "w"
        encoding = "utf-8"
    elif fmt == "docx":
        content = render_docx(data)
        mode = "wb"
        encoding = None
    else:
        raise ValueError(f"Unknown format: {fmt}")

    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if mode == "wb":
        with open(output_path, "wb") as f:
            f.write(content)
    else:
        with open(output_path, "w", encoding=encoding) as f:
            f.write(content)

    return {
        "status": "ok",
        "format": fmt,
        "output": str(output_path.resolve()),
        "size_bytes": output_path.stat().st_size,
    }


if __name__ == "__main__":
    main()
